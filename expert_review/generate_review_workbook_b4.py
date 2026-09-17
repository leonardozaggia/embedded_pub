#!/usr/bin/env python
"""
Generate the Bayley-4 (Step-2) co-author review workbook.

Step 2 validates the unsupervised NLP approach on the Bayley-4: the 39
Aylward-assigned items are clustered with no expert input, and the detected
clusters are compared to the Aylward et al. (2022) reference model. This
workbook lists the *divergent* items (detected cluster ≠ Aylward reference)
and asks two experts to independently re-assign each to one of the five
cognitive domains, exactly as for the Bayley-III review.

Unlike the B3 workbook this one is *pre-filled* with the answers the co-authors
already returned (item_assignments_AD.xlsx / item_assignments_AHildebrandt.xlsx,
consolidated in item_assignments_SOLUTION.xlsx). The generated file is therefore
both an archival record and a template for a future review round.

Output
------
    expert_review/records/coauthor_review_b4.xlsx

Usage
-----
    python expert_review/generate_review_workbook_b4.py
    python expert_review/generate_review_workbook_b4.py --no-prefill      # blank template
    python expert_review/generate_review_workbook_b4.py --output path.xlsx
"""
from __future__ import annotations

import sys as _sys
for _stream in (_sys.stdout, _sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8")
del _stream

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

# ---------------------------------------------------------------------------
# Reuse the B3 generator's styling + parsing helpers (same directory).
# ---------------------------------------------------------------------------
_REPO = Path(__file__).resolve().parents[1]
_sys.path.insert(0, str(Path(__file__).parent))
import generate_review_workbook_b3 as b3   # noqa: E402

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
# The '_no_dr' variant is the one the co-authors actually reviewed (all 16
# reviewed items fall within its 17 divergent items; the alternative variant
# leaves two reviewed items non-divergent).
STEP2_DIR = (
    _REPO / "data/outputs/step2_b4_validation"
    / "all_mpnet_base_v2_kmeans_consensus_k_precomputed_no_dr"
)
ITEMS_CSV      = STEP2_DIR / "items.csv"
COMPARISON_CSV = STEP2_DIR / "05_model_comparison" / "item_comparison_theoretical.csv"

# Existing expert answers (consolidated), as returned by the two co-authors.
DEFAULT_ANSWERS = _REPO / "expert_review/records/b4_raw/item_assignments_SOLUTION.xlsx"
ANSWERS_SHEET   = "Disagreement_Items"

DEFAULT_OUTPUT = _REPO / "expert_review/records/coauthor_review_b4.xlsx"

_FACTOR_ORDER = b3._FACTOR_ORDER


def _norm_title(s) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower()).rstrip(". ")


def _disp(raw) -> str:
    """Display an Aylward/detected factor label; raw text for non-canonical."""
    key = b3.normalize_factor(raw)
    return b3.FACTOR_DISPLAY.get(key, str(raw).strip())


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------

def load_b4_items() -> pd.DataFrame:
    """B4 item text, reusing the B3 materials/admin/scoring parsing."""
    df = pd.read_csv(ITEMS_CSV)
    df["scoring_criteria"] = df["scoring_json"].apply(b3.parse_scoring)
    parsed = df["item_description"].apply(
        lambda d: pd.Series(b3._split_materials_admin(d),
                            index=["materials", "admin_instructions"])
    )
    df = pd.concat([df, parsed], axis=1)
    return df[["item_id", "item_title", "materials", "admin_instructions",
               "scoring_criteria"]]


def load_divergent() -> pd.DataFrame:
    """Items where the detected cluster disagrees with the Aylward reference."""
    cmp = pd.read_csv(COMPARISON_CSV)
    div = cmp[~cmp["agrees"]][
        ["item_id", "reference_factor", "detected_factor"]
    ].copy()
    return div.reset_index(drop=True)


def load_existing_answers(path: Path) -> dict[str, dict[str, str]]:
    """{normalised_title: {'AD': <raw>, 'AH': <raw>}} from the answers workbook."""
    if not path.exists():
        print(f"  [warn] Answers workbook not found ({path}); leaving blank.")
        return {}
    sol = pd.read_excel(path, sheet_name=ANSWERS_SHEET)
    out: dict[str, dict[str, str]] = {}
    for _, r in sol.iterrows():
        if pd.isna(r.get("item_title")):
            continue
        out[_norm_title(r["item_title"])] = {
            "AD": "" if pd.isna(r.get("AD")) else str(r["AD"]).strip(),
            "AH": "" if pd.isna(r.get("AH")) else str(r["AH"]).strip(),
        }
    return out


# ---------------------------------------------------------------------------
# Sheet builder
# ---------------------------------------------------------------------------

def build_discordant_sheet(ws, div: pd.DataFrame, items: pd.DataFrame,
                           answers: dict[str, dict[str, str]]) -> int:
    merged = div.merge(items, on="item_id", how="left")

    # Sort by Aylward reference factor group, then item_id.
    rank = {k: i for i, k in enumerate(_FACTOR_ORDER)}
    merged["_sort"] = merged["reference_factor"].map(
        lambda x: rank.get(b3.normalize_factor(x), 99)
    )
    merged = merged.sort_values(["_sort", "item_id"]).reset_index(drop=True)

    headers = [
        "Item ID",
        "Item Title",
        "Materials",
        "Administration Instructions",
        "Scoring Criteria",
        "Aylward Reference\n(Theoretical)",
        "Detected Assignment\n(Bottom-up NLP)",
        "Expert 1 (A.D.)\nAssignment",
        "Expert 2 (A.H.)\nAssignment",
        "Consensus",
        "Notes",
    ]
    ws.append(headers)

    n_prefilled = 0
    for _, row in merged.iterrows():
        ans = answers.get(_norm_title(row["item_title"]), {})
        ad, ah = ans.get("AD", ""), ans.get("AH", "")
        if ad or ah:
            n_prefilled += 1
        ws.append([
            row["item_id"],
            row["item_title"],
            row.get("materials", ""),
            row.get("admin_instructions", ""),
            row.get("scoring_criteria", ""),
            _disp(row["reference_factor"]),
            _disp(row["detected_factor"]),
            ad,
            ah,
            "",   # Consensus — filled after reconciliation / derived at import
            "",   # Notes
        ])
        r = ws.max_row
        b3.color_factor_cell(ws.cell(r, 6), b3.normalize_factor(row["reference_factor"]))
        b3.color_factor_cell(ws.cell(r, 7), b3.normalize_factor(row["detected_factor"]))

    col_widths = {
        "A": 10, "B": 30, "C": 18, "D": 52, "E": 52,
        "F": 24, "G": 24,
        "H": 18, "I": 18, "J": 18, "K": 35,
    }
    b3.apply_sheet_style(ws, col_widths, freeze_after_col=2)
    return n_prefilled


def build_legend_sheet(ws) -> None:
    ws.append(["Factor", "Abbreviation", "Cell colour"])
    for cell in ws[1]:
        cell.font = b3.Font(bold=True, size=10)
    for factor in _FACTOR_ORDER:
        display = b3.FACTOR_DISPLAY[factor]
        abbrev  = display.split("(")[-1].rstrip(")")
        ws.append([display.split(" (")[0], abbrev, "■"])
        r = ws.max_row
        ws.cell(r, 3).fill = b3.FACTOR_FILL[factor]
        ws.cell(r, 3).font = b3.Font(color="FFFFFF", size=10)
    ws.append([])
    ws.append(["Sheet notes"])
    ws.append([
        "Discordant (Detected vs Aylward):",
        "Bayley-4 items where the unsupervised NLP cluster (Step 2) disagrees "
        "with the Aylward et al. (2022) reference assignment.",
        "Experts: assign each item to one of ATT / WM / GDPS / FS / HOP.",
    ])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 80


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main(output_path: Path, answers_path: Path, prefill: bool) -> None:
    print(f"Output path : {output_path}")
    print(f"Pre-fill    : {prefill}"
          + (f"  (from {answers_path.name})" if prefill else ""))

    items = load_b4_items()
    div   = load_divergent()
    answers = load_existing_answers(answers_path) if prefill else {}
    print(f"Divergent B4 items (detected ≠ Aylward): {len(div)}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    ws1 = wb.create_sheet("Discordant (Det vs Aylward)")
    ws2 = wb.create_sheet("Legend")

    n_prefilled = build_discordant_sheet(ws1, div, items, answers)
    build_legend_sheet(ws2)
    wb.save(output_path)

    print(f"\n{'=' * 60}")
    print(f"  Excel saved: {output_path}")
    print(f"  Divergent items : {len(div)}")
    print(f"  Pre-filled rows : {n_prefilled} (experts' existing answers)")
    print(f"  Unreviewed rows : {len(div) - n_prefilled}")
    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Generate Bayley-4 (Step-2) co-author review workbook."
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT,
                        help=f"Output .xlsx (default: {DEFAULT_OUTPUT.name}).")
    parser.add_argument("--answers", type=Path, default=DEFAULT_ANSWERS,
                        help="Consolidated expert answers workbook "
                             "(default: expert_review/records/b4_raw/item_assignments_SOLUTION.xlsx).")
    parser.add_argument("--no-prefill", dest="prefill", action="store_false",
                        help="Produce a blank template (no existing answers).")
    args = parser.parse_args()
    main(output_path=args.output, answers_path=args.answers, prefill=args.prefill)
