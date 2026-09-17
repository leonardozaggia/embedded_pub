#!/usr/bin/env python
"""
Generate Excel workbook for co-author evaluation of Bayley-III item assignments.

Produces two sheets:

  Sheet 1 – "UC Uncertain (Step 1)"
      Items from Step 1 where the UC similarity exceeded the best EF-factor
      centroid similarity by more than UC_MARGIN_THRESHOLD (default 0.10).
      These items were redirected to their best_ef_factor but their UC
      preference is strong enough to warrant independent expert review.

  Sheet 2 – "Discordant (Step 3 vs Step 1)"
      Items where the bottom-up unsupervised NLP clustering (Step 3) assigned
      a different factor than the theoretical translation pipeline (Step 1).

Experts fill in "Expert 1 Assignment", "Expert 2 Assignment", and "Consensus"
columns for each item.  Each sheet also carries a colour-coded "Step 1
Assignment" (and "Step 3 Assignment" in Sheet 2) for context.

Output
------
    expert_review/records/coauthor_review_b3.xlsx

Usage
-----
    python expert_review/generate_review_workbook_b3.py
    python expert_review/generate_review_workbook_b3.py --centroid-method b3_centroids
    python expert_review/generate_review_workbook_b3.py --output path/to/file.xlsx
"""
from __future__ import annotations

# --- UTF-8 stdout/stderr: this script prints Unicode (arrows, checkmarks)
# which crashes under Windows' default cp1252 encoding; force UTF-8. ---
import sys as _sys
for _stream in (_sys.stdout, _sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8")
del _stream


import argparse
import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------------------
# Repository root resolution
# ---------------------------------------------------------------------------
_REPO = Path(__file__).resolve().parents[1]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
STEP1_DIR = _REPO / "data/outputs/step1_translation"
DEFAULT_CENTROID_METHOD = "mixed_centroids"   # b4_centroids | b3_centroids | mixed_centroids
STEP3_DIR = (
    _REPO
    / "data/outputs/step3_b3_bottomup"
    / "all_mpnet_base_v2_kmeans_consensus_no_dr"
)
ITEMS_CSV = STEP3_DIR / "items.csv"
DEFAULT_OUTPUT = _REPO / "expert_review/records/coauthor_review_b3.xlsx"

UC_MARGIN_THRESHOLD = 0.10

# ---------------------------------------------------------------------------
# Canonical factor metadata
# ---------------------------------------------------------------------------
FACTOR_DISPLAY: dict[str, str] = {
    "attention":                     "Attention (ATT)",
    "working_memory":                "Working Memory (WM)",
    "goal_directed_problem_solving": "Goal-Directed Problem Solving (GDPS)",
    "flexibility_shift":             "Flexibility / Shift (FS)",
    "higher_order_processing":       "Higher-Order Processing (HOP)",
    "uncategorized_cognition":       "Uncategorized (UC)",
}

# Pastel cell background fills per factor
_FACTOR_FILL: dict[str, str] = {
    "attention":                     "F4A7A9",   # pastel red
    "working_memory":                "A8BBDA",   # pastel blue
    "goal_directed_problem_solving": "A8D5B5",   # pastel green
    "flexibility_shift":             "F4C89A",   # pastel orange
    "higher_order_processing":       "C5BCDC",   # pastel purple
    "uncategorized_cognition":       "CCCCCC",   # grey
}
FACTOR_FILL: dict[str, PatternFill] = {
    k: PatternFill("solid", fgColor=v) for k, v in _FACTOR_FILL.items()
}

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
DATA_FONT   = Font(size=10)
WRAP_TOP    = Alignment(wrap_text=True, vertical="top")
_THIN_SIDE  = Side(style="thin", color="CCCCCC")
THIN_BORDER = Border(
    left=_THIN_SIDE, right=_THIN_SIDE,
    top=_THIN_SIDE,  bottom=_THIN_SIDE,
)
YELLOW_FILL = PatternFill("solid", fgColor="FFD966")

# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def normalize_factor(name: str) -> str:
    """Map any factor label variant to a canonical snake_case key."""
    normed = re.sub(r"[\s\-]+", "_", str(name).strip().lower())
    normed = normed.replace("higher_order_processes", "higher_order_processing")
    return normed


def display_factor(factor_key: str) -> str:
    return FACTOR_DISPLAY.get(factor_key, factor_key.replace("_", " ").title())


def parse_scoring(scoring_json_str: str) -> str:
    """Convert JSON scoring string into readable multiline text."""
    try:
        criteria = json.loads(str(scoring_json_str))
        parts = [
            f"[{c['points']} pt] {c['criteria'].rstrip(' -').strip()}"
            for c in criteria
        ]
        return "\n".join(parts)
    except Exception:
        return str(scoring_json_str)


def color_factor_cell(cell, factor_key: str) -> None:
    fill = FACTOR_FILL.get(factor_key)
    if fill:
        cell.fill = fill


def apply_sheet_style(
    ws,
    col_widths: dict[str, float],
    freeze_after_col: int = 2,
) -> None:
    """Style header row, set column widths, freeze panes, wrap all cells."""
    # Header row
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP_TOP
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 40

    # Data rows
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
            cell.border = THIN_BORDER
            cell.font = DATA_FONT

    # Column widths
    from openpyxl.utils import get_column_letter
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Freeze: first N columns + header row
    freeze_cell = ws.cell(row=2, column=freeze_after_col + 1).coordinate
    ws.freeze_panes = freeze_cell


# ---------------------------------------------------------------------------
# Data loaders
# ---------------------------------------------------------------------------

def _split_materials_admin(desc: str) -> tuple[str, str]:
    """
    Split 'Materials: <X> <Admin text>' at the first space+capital-letter
    boundary that does NOT start the word 'Book' (which is part of material
    names like 'Stimulus Book').  If a match starts 'Book', skip it and
    continue searching from the character after that match.
    """
    m = re.match(r"^Materials:\s*(.*)", str(desc), re.DOTALL)
    if not m:
        return "None", str(desc).strip()
    content = m.group(1)
    search_from = 0
    while True:
        split_m = re.search(r" ([A-Z])", content[search_from:])
        if not split_m:
            break
        abs_start = search_from + split_m.start()
        abs_end   = search_from + split_m.end()
        word_rest = re.match(r"[A-Za-z]+", content[abs_end - 1:])
        word = word_rest.group(0) if word_rest else split_m.group(1)
        if word == "Book":
            search_from = abs_end
            continue
        materials = content[:abs_start].strip()
        admin     = content[abs_end - 1:]
        return materials or "None", admin.strip()
    return content.strip() or "None", ""


def load_items() -> pd.DataFrame:
    """Load item metadata from items.csv, parse materials / instructions."""
    df = pd.read_csv(ITEMS_CSV)
    df["scoring_criteria"] = df["scoring_json"].apply(parse_scoring)

    parsed = df["item_description"].apply(
        lambda d: pd.Series(_split_materials_admin(d),
                            index=["materials", "admin_instructions"])
    )
    df = pd.concat([df, parsed], axis=1)

    return df[["item_id", "item_title", "materials", "admin_instructions", "scoring_criteria"]]


def load_step1_assignments(centroid_method: str) -> pd.DataFrame:
    """
    Returns all 91 B3 items with their Step 1 theoretical factor assignment.

    Also carries uc_margin and sim_* columns for unpaired (centroid-assigned)
    items; paired items have NaN in those columns.
    """
    paired_csv   = STEP1_DIR / "b3_paired_assignments.csv"
    unpaired_csv = STEP1_DIR / centroid_method / "all_unpaired_b3_assignments.csv"

    if not paired_csv.exists():
        raise FileNotFoundError(f"Paired assignments not found: {paired_csv}")
    if not unpaired_csv.exists():
        raise FileNotFoundError(f"Unpaired assignments not found: {unpaired_csv}")

    paired   = pd.read_csv(paired_csv)
    unpaired = pd.read_csv(unpaired_csv)

    # Paired items: factor is already final (no UC redirection possible)
    paired = paired.rename(columns={
        "assigned_factor":   "step1_factor",
        "pairing_similarity": "pairing_sim",
    })
    paired["step1_source"] = "cascade_pairing"

    # Unpaired: redirect UC → best_ef_factor
    def _resolve(row):
        if row["assigned_factor"] == "uncategorized_cognition":
            return row["best_ef_factor"]
        return row["assigned_factor"]

    unpaired = unpaired.copy()
    unpaired["step1_factor"] = unpaired.apply(_resolve, axis=1)
    unpaired["step1_source"] = unpaired["assigned_factor"].apply(
        lambda x: "uc_redirected" if x == "uncategorized_cognition" else "centroid_loo"
    )

    sim_cols = [c for c in unpaired.columns if c.startswith("sim_")]
    shared_base = ["item_id", "step1_factor", "step1_source"]

    paired_out   = paired[shared_base].copy()
    paired_out["uc_margin"] = float("nan")
    for c in sim_cols:
        paired_out[c] = float("nan")

    unpaired_out = unpaired[shared_base + ["uc_margin"] + sim_cols].copy()

    return pd.concat([paired_out, unpaired_out], ignore_index=True)


def load_step3_assignments() -> pd.DataFrame:
    """Returns item_id → step3_factor (normalized canonical key)."""
    clust_csv  = STEP3_DIR / "03_clustering" / "cluster_assignments.csv"
    labels_csv = STEP3_DIR / "04_labeling" / "cluster_concept_labels.csv"

    if not clust_csv.exists():
        raise FileNotFoundError(f"Step 3 cluster assignments not found: {clust_csv}")
    if not labels_csv.exists():
        raise FileNotFoundError(f"Step 3 cluster labels not found: {labels_csv}")

    clust  = pd.read_csv(clust_csv)
    labels = pd.read_csv(labels_csv)

    id_to_name      = dict(zip(labels["cluster_id"], labels["top_concept"]))
    clust["step3_factor"] = clust["cluster"].map(id_to_name).apply(normalize_factor)

    return clust[["item_id", "step3_factor"]]


# ---------------------------------------------------------------------------
# Sheet 1 — UC Uncertain items
# ---------------------------------------------------------------------------

def build_uc_uncertain_sheet(ws, step1: pd.DataFrame, items: pd.DataFrame) -> int:
    """
    Populate ws with items where uc_margin > UC_MARGIN_THRESHOLD.

    Returns the number of items written.
    """
    uc_rows = step1[
        step1["uc_margin"].notna() & (step1["uc_margin"] > UC_MARGIN_THRESHOLD)
    ].copy()
    uc_rows = uc_rows.merge(items, on="item_id")
    uc_rows = uc_rows.sort_values("uc_margin", ascending=False).reset_index(drop=True)

    sim_cols    = ["sim_ATT", "sim_WM", "sim_GDPS", "sim_FS", "sim_HOP", "sim_UC"]
    sim_headers = ["Sim ATT", "Sim WM", "Sim GDPS", "Sim FS", "Sim HOP", "Sim UC"]

    headers = [
        "Item ID",
        "Item Title",
        "Materials",
        "Administration Instructions",
        "Scoring Criteria",
        "Step 1 Assignment\n(Theoretical)",
        f"UC Margin\n(UC sim − best EF sim)\nthreshold > {UC_MARGIN_THRESHOLD}",
        *sim_headers,
        "Expert 1\nAssignment",
        "Expert 2\nAssignment",
        "Consensus",
        "Notes",
    ]
    ws.append(headers)

    for _, row in uc_rows.iterrows():
        sim_vals = [
            round(float(row[c]), 4) if pd.notna(row.get(c)) else ""
            for c in sim_cols
        ]
        data = [
            row["item_id"],
            row["item_title"],
            row["materials"],
            row["admin_instructions"],
            row["scoring_criteria"],
            display_factor(row["step1_factor"]),
            round(float(row["uc_margin"]), 4),
            *sim_vals,
            "",  # Expert 1
            "",  # Expert 2
            "",  # Consensus
            "",  # Notes
        ]
        ws.append(data)
        r = ws.max_row
        color_factor_cell(ws.cell(r, 6), row["step1_factor"])
        ws.cell(r, 7).fill = YELLOW_FILL   # highlight uc_margin

    col_widths = {
        "A": 10, "B": 30, "C": 18, "D": 52, "E": 52,
        "F": 26, "G": 14,
        "H": 10, "I": 10, "J": 10, "K": 10, "L": 10, "M": 10,
        "N": 20, "O": 20, "P": 20, "Q": 35,
    }
    apply_sheet_style(ws, col_widths, freeze_after_col=2)
    return len(uc_rows)


# ---------------------------------------------------------------------------
# Sheet 2 — Discordant items (Step 3 ≠ Step 1)
# ---------------------------------------------------------------------------

_FACTOR_ORDER = [
    "attention",
    "working_memory",
    "goal_directed_problem_solving",
    "flexibility_shift",
    "higher_order_processing",
]


def build_discordant_sheet(
    ws,
    step1: pd.DataFrame,
    step3: pd.DataFrame,
    items: pd.DataFrame,
) -> int:
    """
    Populate ws with items where Step 3 factor ≠ Step 1 factor.

    Returns the number of items written.
    """
    merged = step3.merge(
        step1[["item_id", "step1_factor", "step1_source"]],
        on="item_id",
        how="left",
    )
    discordant = merged[merged["step3_factor"] != merged["step1_factor"]].copy()
    discordant = discordant.merge(items, on="item_id")

    # Sort by Step 1 factor group, then item_id
    factor_rank = {k: i for i, k in enumerate(_FACTOR_ORDER)}
    discordant["_sort"] = discordant["step1_factor"].map(factor_rank).fillna(99)
    discordant = (
        discordant.sort_values(["_sort", "item_id"])
        .drop(columns="_sort")
        .reset_index(drop=True)
    )

    headers = [
        "Item ID",
        "Item Title",
        "Materials",
        "Administration Instructions",
        "Scoring Criteria",
        "Step 1 Assignment\n(Theoretical)",
        "Step 3 Assignment\n(Bottom-up NLP)",
        "Assignment Source\n(Step 1)",
        "Expert 1\nAssignment",
        "Expert 2\nAssignment",
        "Consensus",
        "Notes",
    ]
    ws.append(headers)

    for _, row in discordant.iterrows():
        data = [
            row["item_id"],
            row["item_title"],
            row["materials"],
            row["admin_instructions"],
            row["scoring_criteria"],
            display_factor(row["step1_factor"]),
            display_factor(row["step3_factor"]),
            row.get("step1_source", ""),
            "",  # Expert 1
            "",  # Expert 2
            "",  # Consensus
            "",  # Notes
        ]
        ws.append(data)
        r = ws.max_row
        color_factor_cell(ws.cell(r, 6), row["step1_factor"])
        color_factor_cell(ws.cell(r, 7), row["step3_factor"])

    col_widths = {
        "A": 10, "B": 30, "C": 18, "D": 52, "E": 52,
        "F": 26, "G": 26, "H": 20,
        "I": 20, "J": 20, "K": 20, "L": 35,
    }
    apply_sheet_style(ws, col_widths, freeze_after_col=2)
    return len(discordant)


# ---------------------------------------------------------------------------
# Legend sheet
# ---------------------------------------------------------------------------

def build_legend_sheet(ws, centroid_method: str) -> None:
    ws.append(["Factor", "Abbreviation", "Cell colour"])
    ws.row_dimensions[1].height = 20
    for cell in ws[1]:
        cell.font = Font(bold=True, size=10)

    for factor in _FACTOR_ORDER:
        display = FACTOR_DISPLAY[factor]
        abbrev  = display.split("(")[-1].rstrip(")")
        ws.append([display.split(" (")[0], abbrev, "■"])
        r = ws.max_row
        ws.cell(r, 3).fill = FACTOR_FILL[factor]
        ws.cell(r, 3).font = Font(color="FFFFFF", size=10)

    ws.append([])
    ws.append(["Sheet notes"])
    ws.append([
        "Sheet 1 (UC Uncertain):",
        f"Items where UC similarity exceeded the best EF-factor centroid by > {UC_MARGIN_THRESHOLD}.",
        "Experts: assign each item to one of ATT / WM / GDPS / FS / HOP.",
    ])
    ws.append([
        "Sheet 2 (Discordant):",
        "Items where the bottom-up NLP clustering (Step 3) disagrees with the theoretical assignment (Step 1).",
        "Experts: assign each item to one of ATT / WM / GDPS / FS / HOP.",
    ])
    ws.append([
        "Centroid method used for Step 1:",
        centroid_method,
    ])

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 80


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main(centroid_method: str, output_path: Path) -> None:
    print(f"Centroid method : {centroid_method}")
    print(f"Output path     : {output_path}\n")

    print("Loading data …")
    items = load_items()
    step1 = load_step1_assignments(centroid_method)
    step3 = load_step3_assignments()

    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)   # remove default empty sheet

    ws1 = wb.create_sheet("UC Uncertain (Step 1)")
    ws2 = wb.create_sheet("Discordant (Step 3 vs Step 1)")
    ws3 = wb.create_sheet("Legend")

    print("Building Sheet 1 — UC Uncertain …")
    n1 = build_uc_uncertain_sheet(ws1, step1, items)

    print("Building Sheet 2 — Discordant …")
    n2 = build_discordant_sheet(ws2, step1, step3, items)

    build_legend_sheet(ws3, centroid_method)

    wb.save(output_path)

    print(f"\n{'=' * 60}")
    print(f"  Excel saved: {output_path}")
    print(f"  Sheet 1 — UC Uncertain:  {n1} items  (uc_margin > {UC_MARGIN_THRESHOLD})")
    print(f"  Sheet 2 — Discordant:    {n2} items  (Step 3 ≠ Step 1)")
    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Generate Bayley-III co-author review Excel."
    )
    parser.add_argument(
        "--centroid-method",
        choices=["b4_centroids", "b3_centroids", "mixed_centroids"],
        default=DEFAULT_CENTROID_METHOD,
        help="Which Step 1 centroid method to use (default: mixed_centroids).",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Output .xlsx path (default: expert_review/records/coauthor_review_b3.xlsx).",
    )
    args = parser.parse_args()
    main(centroid_method=args.centroid_method, output_path=args.output)
